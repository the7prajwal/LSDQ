"""
Multi-format document loaders for RAG.
Supports: .pdf, .txt, .docx, .pptx, .xlsx

Each loader returns a list of "segments": small dicts of
    {"text": str, "location": str}
where "location" is a human-readable pointer (page 3, slide 5, paragraph 12, etc.)
so retrieved chunks can be traced back to exactly where they came from.

Requirements:
    pip install pypdf python-docx python-pptx openpyxl
"""

import os


def load_txt(path: str) -> list:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        lines = f.readlines()

    segments = []
    para_lines = []
    para_start = 1
    for i, line in enumerate(lines, start=1):
        if line.strip() == "":
            if para_lines:
                text = "".join(para_lines).strip()
                segments.append({
                    "text": text,
                    "location": f"lines {para_start}-{i - 1}",
                })
                para_lines = []
            para_start = i + 1
        else:
            para_lines.append(line)
    if para_lines:
        text = "".join(para_lines).strip()
        segments.append({
            "text": text,
            "location": f"lines {para_start}-{len(lines)}",
        })
    return segments


def load_pdf(path: str) -> list:
    from pypdf import PdfReader
    reader = PdfReader(path)
    segments = []
    for page_num, page in enumerate(reader.pages, start=1):
        text = page.extract_text()
        if text and text.strip():
            segments.append({
                "text": text.strip(),
                "location": f"page {page_num}",
            })
    return segments


def load_docx(path: str) -> list:
    import docx
    doc = docx.Document(path)
    segments = []

    for i, p in enumerate(doc.paragraphs, start=1):
        if p.text.strip():
            segments.append({
                "text": p.text.strip(),
                "location": f"paragraph {i}",
            })

    for t_idx, table in enumerate(doc.tables, start=1):
        for r_idx, row in enumerate(table.rows, start=1):
            row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_text:
                segments.append({
                    "text": " | ".join(row_text),
                    "location": f"table {t_idx}, row {r_idx}",
                })

    return segments


def load_pptx(path: str) -> list:
    from pptx import Presentation
    prs = Presentation(path)
    segments = []

    for slide_num, slide in enumerate(prs.slides, start=1):
        slide_text = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in paragraph.runs)
                    if text.strip():
                        slide_text.append(text.strip())
        if slide_text:
            segments.append({
                "text": "\n".join(slide_text),
                "location": f"slide {slide_num}",
            })

    return segments


def load_xlsx(path: str) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True)
    segments = []

    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        for row_num, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            row_text = [str(cell) for cell in row if cell is not None]
            if row_text:
                segments.append({
                    "text": " | ".join(row_text),
                    "location": f"sheet '{sheet_name}', row {row_num}",
                })

    return segments


LOADERS = {
    ".txt": load_txt,
    ".md": load_txt,
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".pptx": load_pptx,
    ".xlsx": load_xlsx,
}


def load_document(path: str) -> list:
    ext = os.path.splitext(path)[1].lower()
    loader = LOADERS.get(ext)
    if loader is None:
        raise ValueError(f"Unsupported file type: {ext} ({path})")
    return loader(path)


def load_folder(folder_path: str) -> dict:
    results = {}
    for root, _, files in os.walk(folder_path):
        for filename in files:
            ext = os.path.splitext(filename)[1].lower()
            if ext in LOADERS:
                full_path = os.path.join(root, filename)
                try:
                    segments = load_document(full_path)
                    if segments:
                        results[full_path] = segments
                except Exception as e:
                    print(f"Failed to load {full_path}: {e}")
    return results
